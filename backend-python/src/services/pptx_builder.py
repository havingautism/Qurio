"""
HTML pagination and PPTX generation helpers for ppt_generator tool.
"""

from __future__ import annotations

import asyncio
import base64
import json
import io
import re
from typing import Any

import httpx
from bs4 import BeautifulSoup

from .pptx_schema import build_pptx_error

BLOCK_WEIGHTS: dict[str, int] = {
    "heading": 60,
    "paragraph": 20,
    "bullets": 30,
    "table": 120,
    "image": 160,
}


def _sanitize_html(raw_html: Any) -> str:
    html = str(raw_html or "").strip()
    if not html:
        return ""
    html = re.sub(r"<\s*script[\s\S]*?<\s*/\s*script\s*>", "", html, flags=re.IGNORECASE)
    html = re.sub(r"<\s*(iframe|object|embed|link|base)\b[\s\S]*?<\s*/\s*\1\s*>", "", html, flags=re.IGNORECASE)
    html = re.sub(r"<\s*(iframe|object|embed|link|base)\b[^>]*?/?>", "", html, flags=re.IGNORECASE)
    html = re.sub(r"\son\w+\s*=\s*\"[^\"]*\"", "", html, flags=re.IGNORECASE)
    html = re.sub(r"\son\w+\s*=\s*'[^']*'", "", html, flags=re.IGNORECASE)
    html = re.sub(r"\son\w+\s*=\s*[^\s>]+", "", html, flags=re.IGNORECASE)
    html = re.sub(r"javascript\s*:", "", html, flags=re.IGNORECASE)
    return html


def _clean_text(value: str) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _is_page_break_tag(tag: Any) -> bool:
    if not getattr(tag, "name", None):
        return False
    if tag.attrs.get("data-page-break"):
        return True
    style = str(tag.attrs.get("style") or "").lower()
    return "page-break-before" in style or "break-before" in style


def _extract_blocks_from_container(container: Any, respect_page_break: bool) -> list[dict[str, Any]]:
    blocks: list[dict[str, Any]] = []
    for child in getattr(container, "children", []):
        name = getattr(child, "name", None)
        if not name:
            text = _clean_text(str(child))
            if text:
                blocks.append({"type": "paragraph", "text": text})
            continue

        tag_name = str(name).lower()

        if tag_name in {"style", "script", "noscript"}:
            continue

        if respect_page_break and _is_page_break_tag(child):
            blocks.append({"type": "page_break"})
            continue

        if tag_name in {"h1", "h2", "h3", "h4"}:
            text = _clean_text(child.get_text(" ", strip=True))
            if text:
                blocks.append({"type": "heading", "level": int(tag_name[1]), "text": text})
            continue

        if tag_name in {"p", "blockquote"}:
            text = _clean_text(child.get_text(" ", strip=True))
            if text:
                blocks.append({"type": "paragraph", "text": text})
            continue

        if tag_name in {"div", "section", "article", "main"}:
            # Avoid flattening whole layout wrappers into one giant paragraph.
            nested = _extract_blocks_from_container(child, respect_page_break=respect_page_break)
            if nested:
                blocks.extend(nested)
                continue
            text = _clean_text(child.get_text(" ", strip=True))
            if text:
                blocks.append({"type": "paragraph", "text": text})
            continue

        if tag_name in {"ul", "ol"}:
            items = [_clean_text(li.get_text(" ", strip=True)) for li in child.find_all("li", recursive=False)]
            items = [item for item in items if item]
            if items:
                blocks.append({"type": "bullets", "ordered": tag_name == "ol", "items": items})
            continue

        if tag_name == "table":
            rows: list[list[str]] = []
            for tr in child.find_all("tr"):
                cells = [_clean_text(td.get_text(" ", strip=True)) for td in tr.find_all(["th", "td"])]
                if any(cells):
                    rows.append(cells)
            if rows:
                blocks.append({"type": "table", "rows": rows})
            continue

        if tag_name == "img":
            src = str(child.attrs.get("src") or "").strip()
            if src:
                blocks.append({"type": "image", "src": src})
            continue

        nested = _extract_blocks_from_container(child, respect_page_break=respect_page_break)
        blocks.extend(nested)

    return blocks


def _split_text_chunks(text: str, limit: int) -> list[str]:
    clean = _clean_text(text)
    if len(clean) <= limit:
        return [clean]

    sentences = [seg.strip() for seg in re.split(r"(?<=[.!?。！？])\s+", clean) if seg.strip()]
    if len(sentences) <= 1:
        words = clean.split(" ")
        chunks: list[str] = []
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            if len(candidate) <= limit:
                current = candidate
            else:
                if current:
                    chunks.append(current)
                current = word
        if current:
            chunks.append(current)
        return chunks

    chunks = []
    current = ""
    for sentence in sentences:
        candidate = f"{current} {sentence}".strip()
        if len(candidate) <= limit:
            current = candidate
        else:
            if current:
                chunks.append(current)
            current = sentence
    if current:
        chunks.append(current)
    return chunks


def _block_cost(block: dict[str, Any]) -> int:
    block_type = block.get("type")
    if block_type == "heading":
        return BLOCK_WEIGHTS["heading"] + len(str(block.get("text") or ""))
    if block_type == "paragraph":
        return BLOCK_WEIGHTS["paragraph"] + len(str(block.get("text") or ""))
    if block_type == "bullets":
        text_len = sum(len(str(item)) for item in block.get("items", []))
        return BLOCK_WEIGHTS["bullets"] + text_len
    if block_type == "table":
        text_len = sum(len(str(cell)) for row in block.get("rows", []) for cell in row)
        return BLOCK_WEIGHTS["table"] + text_len
    if block_type == "image":
        return BLOCK_WEIGHTS["image"]
    return 0


def _split_oversized_block(block: dict[str, Any], max_chars_per_slide: int) -> list[dict[str, Any]]:
    block_type = block.get("type")
    if block_type in {"heading", "paragraph"}:
        parts = _split_text_chunks(str(block.get("text") or ""), max(80, max_chars_per_slide - 100))
        if block_type == "heading":
            return [{"type": "heading", "level": int(block.get("level") or 2), "text": part} for part in parts]
        return [{"type": "paragraph", "text": part} for part in parts]

    if block_type == "bullets":
        chunks: list[list[str]] = []
        current: list[str] = []
        current_cost = 0
        for item in block.get("items", []):
            item_text = str(item)
            item_cost = len(item_text) + 20
            if current and current_cost + item_cost > max_chars_per_slide:
                chunks.append(current)
                current = [item_text]
                current_cost = item_cost
            else:
                current.append(item_text)
                current_cost += item_cost
        if current:
            chunks.append(current)
        return [{"type": "bullets", "ordered": bool(block.get("ordered")), "items": chunk} for chunk in chunks]

    if block_type == "table":
        rows = block.get("rows", [])
        if len(rows) <= 2:
            return [block]
        header = rows[0]
        data_rows = rows[1:]
        chunks: list[list[list[str]]] = []
        current = [header]
        current_cost = _block_cost({"type": "table", "rows": [header]})
        for row in data_rows:
            row_cost = sum(len(str(cell)) for cell in row) + 20
            if len(current) > 1 and current_cost + row_cost > max_chars_per_slide:
                chunks.append(current)
                current = [header, row]
                current_cost = _block_cost({"type": "table", "rows": [header, row]})
            else:
                current.append(row)
                current_cost += row_cost
        if current:
            chunks.append(current)
        return [{"type": "table", "rows": chunk} for chunk in chunks]

    return [block]


def _paginate_blocks(
    blocks: list[dict[str, Any]],
    max_chars_per_slide: int,
    respect_page_break: bool,
) -> list[list[dict[str, Any]]]:
    slides: list[list[dict[str, Any]]] = []
    current: list[dict[str, Any]] = []
    current_cost = 0

    def flush() -> None:
        nonlocal current, current_cost
        if current:
            slides.append(current)
            current = []
            current_cost = 0

    for block in blocks:
        if block.get("type") == "page_break" and respect_page_break:
            flush()
            continue

        split_blocks = _split_oversized_block(block, max_chars_per_slide)
        for piece in split_blocks:
            piece_cost = _block_cost(piece)
            if current and current_cost + piece_cost > max_chars_per_slide:
                flush()
            current.append(piece)
            current_cost += piece_cost

    flush()
    return slides


def paginate_html(
    html: str,
    *,
    mode: str = "auto",
    selector: str = "",
    max_chars_per_slide: int = 900,
    respect_page_break: bool = True,
) -> list[list[dict[str, Any]]]:
    safe_html = _sanitize_html(html)
    if not safe_html:
        return []

    soup = BeautifulSoup(safe_html, "html.parser")
    body = soup.body or soup

    auto_selector = ""
    if mode == "auto":
        for candidate in (".slide", "[data-slide]", "section.slide"):
            try:
                matched = body.select(candidate)
            except Exception:
                matched = []
            if len(matched) >= 2:
                auto_selector = candidate
                break

    effective_mode = "selector" if (mode == "selector" and selector) or auto_selector else mode
    effective_selector = selector if selector else auto_selector

    if effective_mode == "selector" and effective_selector:
        selected = body.select(effective_selector)
        blocks: list[dict[str, Any]] = []
        for index, node in enumerate(selected):
            node_blocks = _extract_blocks_from_container(node, respect_page_break=False)
            blocks.extend(node_blocks)
            if index != len(selected) - 1:
                blocks.append({"type": "page_break"})
    else:
        blocks = _extract_blocks_from_container(body, respect_page_break=respect_page_break)

    return _paginate_blocks(
        blocks,
        max_chars_per_slide=max_chars_per_slide,
        respect_page_break=respect_page_break,
    )


def apply_slide_overrides(
    slides: list[list[dict[str, Any]]],
    slide_overrides: dict[int, dict[str, Any]] | None = None,
) -> list[list[dict[str, Any]]]:
    if not slide_overrides:
        return slides
    out: list[list[dict[str, Any]]] = []
    for index, slide in enumerate(slides, start=1):
        override = slide_overrides.get(index) if isinstance(slide_overrides, dict) else None
        if not isinstance(override, dict):
            out.append(slide)
            continue
        title = str(override.get("title") or "").strip()
        if not title:
            out.append(slide)
            continue
        if slide and slide[0].get("type") == "heading":
            next_slide = [dict(slide[0]), *slide[1:]]
            next_slide[0]["text"] = title
            out.append(next_slide)
        else:
            out.append([{"type": "heading", "level": 2, "text": title}, *slide])
    return out


def _normalize_hex_color(value: str, default: str) -> str:
    candidate = str(value or "").strip().lstrip("#")
    if re.fullmatch(r"[0-9a-fA-F]{6}", candidate):
        return candidate.upper()
    return default


def _parse_data_image(image_src: str) -> io.BytesIO | None:
    try:
        _, payload = image_src.split(",", 1)
        raw = base64.b64decode(payload)
        return io.BytesIO(raw)
    except Exception:
        return None


async def _fetch_remote_image_bytes(image_url: str, max_bytes: int = 8 * 1024 * 1024) -> io.BytesIO | None:
    url = str(image_url or "").strip()
    if not url.startswith(("http://", "https://")):
        return None
    try:
        timeout = httpx.Timeout(8.0, connect=4.0)
        async with httpx.AsyncClient(timeout=timeout, follow_redirects=True) as client:
            resp = await client.get(url)
            if resp.status_code >= 400:
                return None
            data = resp.content or b""
            if not data or len(data) > max_bytes:
                return None
            return io.BytesIO(data)
    except Exception:
        return None


async def _load_image_stream(image_src: str) -> io.BytesIO | None:
    src = str(image_src or "").strip()
    if not src:
        return None
    if src.startswith("data:image/") and ";base64," in src:
        return _parse_data_image(src)
    return await _fetch_remote_image_bytes(src)


def _selector_candidate_order(explicit_selector: str = "") -> list[str]:
    candidates: list[str] = []
    if explicit_selector:
        candidates.append(explicit_selector)
    for c in (".slide", "[data-slide]", "section.slide"):
        if c not in candidates:
            candidates.append(c)
    return candidates


def _extract_stylesheet_tags(html: str, max_style_tags: int = 8) -> str:
    if not html:
        return ""
    soup = BeautifulSoup(html, "html.parser")
    style_chunks: list[str] = []
    for tag in soup.find_all("style")[:max_style_tags]:
        style_chunks.append(str(tag))
    for tag in soup.find_all("link"):
        rel = " ".join(tag.get("rel") or [])
        if "stylesheet" in rel.lower():
            style_chunks.append(str(tag))
    return "".join(style_chunks)


def _select_slide_nodes(html: str, explicit_selector: str = "") -> tuple[list[str], str]:
    if not html:
        return [], ""
    soup = BeautifulSoup(html, "html.parser")
    body = soup.body or soup

    for candidate in _selector_candidate_order(explicit_selector):
        try:
            selected = body.select(candidate)
        except Exception:
            selected = []
        if not selected:
            continue
        if explicit_selector and candidate == explicit_selector:
            return [str(node) for node in selected], candidate
        if len(selected) >= 2:
            return [str(node) for node in selected], candidate

    return [], ""


def _normalize_slide_html_units(request_payload: dict[str, Any]) -> tuple[list[str], str, str]:
    raw_slides = request_payload.get("slides_html")
    if isinstance(raw_slides, list):
        cleaned_units = [_sanitize_html(item) for item in raw_slides]
        cleaned_units = [unit for unit in cleaned_units if unit]
        if cleaned_units:
            shared_styles = _extract_stylesheet_tags("".join(cleaned_units))
            return cleaned_units, "slides_html", shared_styles

    html = _sanitize_html(request_payload.get("html") or "")
    if not html:
        return [], "empty", ""

    explicit_selector = str(request_payload.get("paginate", {}).get("selector") or "")
    selected_units, source = _select_slide_nodes(html, explicit_selector=explicit_selector)
    if selected_units:
        return selected_units, source or "selector", _extract_stylesheet_tags(html)

    return [html], "single_html", _extract_stylesheet_tags(html)


def collect_pptx_input_qa_issues(raw_html: str) -> list[str]:
    raw = str(raw_html or "").strip()
    if not raw:
        return ["empty_html"]

    issues: list[str] = []
    raw_soup = BeautifulSoup(raw, "html.parser")
    soup = BeautifulSoup(_sanitize_html(raw), "html.parser")
    style_tags = soup.find_all("style")
    stylesheet_links = []
    for tag in raw_soup.find_all("link"):
        rel = " ".join(tag.get("rel") or [])
        if "stylesheet" in rel.lower():
            stylesheet_links.append(tag)

    class_nodes = [node for node in soup.find_all(True) if node.get("class")]
    slide_like_nodes = []
    try:
        slide_like_nodes = list((soup.body or soup).select(".slide"))
    except Exception:
        slide_like_nodes = []

    if stylesheet_links and not style_tags:
        issues.append("external_css_only")
    if class_nodes and not style_tags:
        issues.append("missing_inline_styles")
    if len(slide_like_nodes) <= 1 and len(class_nodes) > 4:
        issues.append("single_html_preview_may_not_match_export_pagination")

    deduped: list[str] = []
    for item in issues:
        if item not in deduped:
            deduped.append(item)
    return deduped


def _truncate_issue(message: str, max_len: int = 220) -> str:
    text = str(message or "").strip().replace("\n", " ")
    if len(text) <= max_len:
        return text
    return text[: max_len - 3].rstrip() + "..."


def _build_html_deck_preview(slide_html_units: list[str], shared_styles: str = "") -> str:
    if not slide_html_units:
        return ""
    docs: list[str] = []
    max_pages = min(len(slide_html_units), 30)
    for raw_slide in slide_html_units[:max_pages]:
        safe_slide = _sanitize_html(raw_slide)
        if not safe_slide:
            continue
        srcdoc = (
            "<!doctype html><html><head><meta charset='utf-8'/>"
            "<meta name='viewport' content='width=device-width, initial-scale=1'/>"
            "<style>html,body{margin:0;padding:0;width:100%;height:100%;overflow:hidden;background:#0b0f17;}"
            "*{box-sizing:border-box;max-width:100%;} img,video,canvas,svg{max-width:100%;height:auto;}"
            ".qurio-slide-root{width:100%;height:100%;overflow:auto;}"
            ".qurio-slide-root .slide{opacity:1!important;display:flex!important;position:relative!important;left:auto!important;top:auto!important;transform:none!important;}"
            ".qurio-slide-root .slide.active{opacity:1!important;}"
            ".qurio-slide-root .dots-container,.qurio-slide-root .navigation,.qurio-slide-root .nav,.qurio-slide-root [class*='nav-dot'],"
            ".qurio-slide-root .swiper-pagination,.qurio-slide-root .swiper-button-prev,.qurio-slide-root .swiper-button-next,"
            ".qurio-slide-root .prev,.qurio-slide-root .next,.qurio-slide-root [class*='prev'],.qurio-slide-root [class*='next'],"
            ".qurio-slide-root [id*='prev'],.qurio-slide-root [id*='next'],.qurio-slide-root [aria-label*='Prev'],.qurio-slide-root [aria-label*='Next']{display:none!important;}"
            "</style>"
            + shared_styles
            + "</head><body><div class='qurio-slide-root'>"
            + safe_slide
            + "</div></body></html>"
        )
        docs.append(srcdoc)

    if not docs:
        return ""

    docs_json = json.dumps(docs, ensure_ascii=False)
    return (
        "<div class='qurio-ppt-root'>"
        "<style>"
        ".qurio-ppt-root{padding:10px;color:#eaf0ff;font-family:ui-sans-serif,system-ui,sans-serif;}"
        ".qurio-ppt-toolbar{display:flex;align-items:center;justify-content:flex-end;gap:8px;margin-bottom:8px;}"
        ".qurio-ppt-actions{display:flex;gap:6px;}"
        ".qurio-ppt-btn{font-size:11px;padding:4px 8px;border-radius:999px;border:1px solid #334155;background:#111827;color:#dbe7ff;cursor:pointer;}"
        ".qurio-ppt-btn:disabled{opacity:.45;cursor:not-allowed;}"
        ".qurio-ppt-frame{width:100%;height:680px;border:1px solid #2a3140;border-radius:12px;background:#0b0f17;}"
        ".qurio-ppt-dots{margin-top:8px;display:flex;gap:6px;flex-wrap:wrap}"
        ".qurio-ppt-dot{font-size:11px;line-height:1;padding:5px 8px;border-radius:999px;border:1px solid #334155;background:#111827;color:#dbe7ff;cursor:pointer;}"
        ".qurio-ppt-dot.active{border-color:#7aa2ff;background:#1e293b;color:#ffffff;}"
        "@media (prefers-color-scheme: light){"
        ".qurio-ppt-root{color:#1f2937;}"
        ".qurio-ppt-btn{border-color:#cbd5e1;background:#f8fafc;color:#334155;}"
        ".qurio-ppt-frame{border-color:#d7dee8;background:#f8fafc;}"
        ".qurio-ppt-dot{border-color:#cbd5e1;background:#f8fafc;color:#334155;}"
        ".qurio-ppt-dot.active{border-color:#5b79b6;background:#eaf1ff;color:#1e3a8a;}"
        "}"
        "</style>"
        "<div class='qurio-ppt-toolbar'>"
        "<div class='qurio-ppt-actions'><button type='button' class='qurio-ppt-btn' id='qurioPrev'>Prev</button>"
        "<button type='button' class='qurio-ppt-btn' id='qurioNext'>Next</button></div></div>"
        "<iframe id='qurioPptFrame' class='qurio-ppt-frame' sandbox='allow-scripts' referrerpolicy='no-referrer'></iframe>"
        "<div id='qurioPptDots' class='qurio-ppt-dots'></div>"
        "<script>"
        f"const qurioDocs={docs_json};"
        "let qurioIndex=0;"
        "const qurioFrame=document.getElementById('qurioPptFrame');"
        "const qurioPrev=document.getElementById('qurioPrev');"
        "const qurioNext=document.getElementById('qurioNext');"
        "const qurioDots=document.getElementById('qurioPptDots');"
        "function qurioRender(i){"
        " if(!qurioDocs.length){return;}"
        " qurioIndex=Math.max(0,Math.min(qurioDocs.length-1,i));"
        " qurioFrame.srcdoc=qurioDocs[qurioIndex];"
        " qurioPrev.disabled=qurioIndex===0;"
        " qurioNext.disabled=qurioIndex===qurioDocs.length-1;"
        " Array.from(qurioDots.children).forEach((el,idx)=>el.classList.toggle('active',idx===qurioIndex));"
        "}"
        "qurioDocs.forEach((_,idx)=>{"
        " const dot=document.createElement('button');"
        " dot.type='button';"
        " dot.className='qurio-ppt-dot';"
        " dot.textContent=String(idx+1);"
        " dot.addEventListener('click',()=>qurioRender(idx));"
        " qurioDots.appendChild(dot);"
        "});"
        "qurioPrev.addEventListener('click',()=>qurioRender(qurioIndex-1));"
        "qurioNext.addEventListener('click',()=>qurioRender(qurioIndex+1));"
        "qurioRender(0);"
        "</script>"
        "</div>"
    )


def _build_fidelity_slide_document(raw_slide: str, shared_styles: str = "") -> str:
    safe_slide = _sanitize_html(raw_slide)
    return (
        "<!doctype html><html><head><meta charset='utf-8'/>"
        "<meta name='viewport' content='width=device-width, initial-scale=1'/>"
        "<style>"
        "html,body{margin:0;padding:0;width:100%;height:100%;background:transparent;overflow:hidden;}"
        "*{box-sizing:border-box;max-width:100%;}"
        "img,video,canvas,svg{max-width:100%;height:auto;}"
        ".qurio-slide-stage{width:100vw;height:100vh;overflow:hidden;}"
        ".qurio-slide-stage .slide{opacity:1!important;display:flex!important;position:relative!important;left:auto!important;top:auto!important;transform:none!important;}"
        ".qurio-slide-stage .slide.active{opacity:1!important;}"
        ".qurio-slide-stage .dots-container,.qurio-slide-stage .navigation,.qurio-slide-stage .nav,.qurio-slide-stage [class*='nav-dot'],"
        ".qurio-slide-stage .swiper-pagination,.qurio-slide-stage .swiper-button-prev,.qurio-slide-stage .swiper-button-next{display:none!important;}"
        "</style>"
        + shared_styles
        + "</head><body><div class='qurio-slide-stage'>"
        + safe_slide
        + "</div></body></html>"
    )


def _slides_from_html_units(
    units: list[str],
    *,
    max_chars_per_slide: int,
    respect_page_break: bool,
    force_one_slide_per_unit: bool,
) -> list[list[dict[str, Any]]]:
    output: list[list[dict[str, Any]]] = []
    for unit_html in units:
        safe_unit = _sanitize_html(unit_html)
        if not safe_unit:
            continue
        if force_one_slide_per_unit:
            soup = BeautifulSoup(safe_unit, "html.parser")
            body = soup.body or soup
            blocks = _extract_blocks_from_container(body, respect_page_break=False)
            if blocks:
                output.append(blocks)
            continue

        paged = paginate_html(
            safe_unit,
            mode="auto",
            selector="",
            max_chars_per_slide=max_chars_per_slide,
            respect_page_break=respect_page_break,
        )
        output.extend(paged)
    return output


def _estimate_paragraph_height(text: str) -> float:
    lines = max(1, (len(str(text or "")) // 70) + 1)
    return max(0.2, 0.32 * lines + 0.12)


def _estimate_bullets_height(items: list[str]) -> float:
    return max(0.8, len(items) * 0.36 + 0.12)


def _estimate_table_height(rows: list[list[str]]) -> float:
    return max(1.0, len(rows) * 0.35 + 0.2)


def _split_block_for_available_height(
    block: dict[str, Any],
    available_height: float,
) -> tuple[dict[str, Any] | None, dict[str, Any] | None]:
    block_type = block.get("type")
    if block_type == "paragraph":
        text = str(block.get("text") or "").strip()
        if not text:
            return None, None
        line_budget = int((available_height - 0.12) // 0.32)
        if line_budget <= 0:
            return None, block
        char_budget = max(40, line_budget * 70)
        chunks = _split_text_chunks(text, char_budget)
        if not chunks:
            return None, block
        head = {"type": "paragraph", "text": chunks[0]}
        tail_text = " ".join(chunks[1:]).strip()
        tail = {"type": "paragraph", "text": tail_text} if tail_text else None
        return head, tail

    if block_type == "bullets":
        items = [str(i).strip() for i in block.get("items", []) if str(i).strip()]
        if not items:
            return None, None
        max_items = int((available_height - 0.12) // 0.36)
        if max_items <= 0:
            return None, block
        if max_items >= len(items):
            return block, None
        head = {"type": "bullets", "ordered": bool(block.get("ordered")), "items": items[:max_items]}
        tail = {"type": "bullets", "ordered": bool(block.get("ordered")), "items": items[max_items:]}
        return head, tail

    if block_type == "table":
        rows = block.get("rows", [])
        if not rows:
            return None, None
        rows_budget = int((available_height - 0.2) // 0.35)
        if rows_budget <= 1:
            return None, block
        if rows_budget >= len(rows):
            return block, None
        header = rows[0]
        slice_count = max(2, rows_budget)
        head_rows = rows[:slice_count]
        tail_rows = [header, *rows[slice_count:]]
        head = {"type": "table", "rows": head_rows}
        tail = {"type": "table", "rows": tail_rows} if len(tail_rows) > 1 else None
        return head, tail

    return None, block


def _semantic_block_height(block: dict[str, Any]) -> float:
    block_type = block.get("type")
    if block_type == "heading":
        return 0.88
    if block_type == "paragraph":
        return _estimate_paragraph_height(str(block.get("text") or ""))
    if block_type == "bullets":
        return _estimate_bullets_height([str(i) for i in block.get("items", []) if str(i).strip()])
    if block_type == "table":
        return _estimate_table_height(block.get("rows", []))
    if block_type == "image":
        return 2.0
    return 0.4


def _build_semantic_slide_doc(slide_blocks: list[dict[str, Any]], slide_index: int) -> str:
    title = f"Slide {slide_index + 1}"
    fragments: list[str] = []
    for block in slide_blocks:
        block_type = block.get("type")
        if block_type == "heading":
            text = str(block.get("text") or "").strip()
            if text:
                if title.startswith("Slide"):
                    title = text
                fragments.append(f"<h2>{text}</h2>")
        elif block_type == "paragraph":
            text = str(block.get("text") or "").strip()
            if text:
                fragments.append(f"<p>{text}</p>")
        elif block_type == "bullets":
            items = [str(i).strip() for i in block.get("items", []) if str(i).strip()]
            if items:
                items_html = "".join(f"<li>{item}</li>" for item in items)
                tag = "ol" if bool(block.get("ordered")) else "ul"
                fragments.append(f"<{tag}>{items_html}</{tag}>")
        elif block_type == "table":
            rows = block.get("rows", [])
            if rows:
                row_html = []
                for row_index, row in enumerate(rows):
                    cell_tag = "th" if row_index == 0 else "td"
                    cells = "".join(f"<{cell_tag}>{str(cell)}</{cell_tag}>" for cell in row)
                    row_html.append(f"<tr>{cells}</tr>")
                fragments.append(f"<table>{''.join(row_html)}</table>")
        elif block_type == "image":
            src = str(block.get("src") or "").strip()
            if src:
                fragments.append(f"<img src=\"{src}\" alt=\"slide image\" />")

    body = "".join(fragments) or "<p>(No slide content)</p>"
    return (
        "<style>"
        ".semantic-slide{padding:28px;min-height:100vh;box-sizing:border-box;background:#0b0f17;color:#e8edf8;font-family:ui-sans-serif,system-ui,sans-serif;}"
        ".semantic-slide .title{font-size:12px;letter-spacing:.08em;text-transform:uppercase;opacity:.65;margin-bottom:12px;}"
        ".semantic-slide h1,.semantic-slide h2,.semantic-slide h3{margin:0 0 12px;line-height:1.3;color:#f4f7ff;}"
        ".semantic-slide p{margin:0 0 10px;line-height:1.6;}"
        ".semantic-slide ul,.semantic-slide ol{margin:0 0 12px 1.1rem;padding:0;}"
        ".semantic-slide li{margin:0 0 6px;line-height:1.5;}"
        ".semantic-slide table{width:100%;border-collapse:collapse;margin:0 0 12px;background:#111827;border:1px solid #2b3448;}"
        ".semantic-slide th,.semantic-slide td{border:1px solid #2b3448;padding:8px 10px;font-size:13px;text-align:left;}"
        ".semantic-slide th{background:#182235;}"
        ".semantic-slide img{max-width:100%;height:auto;border-radius:10px;border:1px solid #2b3448;display:block;}"
        "</style>"
        "<div class='semantic-slide'>"
        f"<div class='title'>Slide {slide_index + 1}</div>"
        f"<h1>{title}</h1>"
        f"{body}"
        "</div>"
    )


async def _render_fidelity_png_slides(
    html: str,
    selector: str = "",
    width_px: int = 1600,
    height_px: int = 900,
) -> tuple[list[bytes], list[str], str]:
    issues: list[str] = []
    images: list[bytes] = []
    selector_used = selector

    try:
        from playwright.async_api import async_playwright
    except Exception as exc:
        return [], [f"fidelity_unavailable_playwright_missing:{exc}"], selector_used

    try:
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=True)
            page = await browser.new_page(viewport={"width": width_px, "height": height_px})
            await page.set_content(html, wait_until="networkidle", timeout=20000)
            await page.wait_for_timeout(400)

            count = 0
            selected_candidate = selector
            for candidate in _selector_candidate_order(selector):
                try:
                    c = await page.locator(candidate).count()
                except Exception:
                    c = 0
                if c > 0:
                    selected_candidate = candidate
                    count = c
                    break

            selector_used = selected_candidate
            if count <= 0:
                png = await page.locator("body").screenshot(type="png")
                images.append(png)
                await browser.close()
                return images, issues, selector_used

            for index in range(count):
                await page.evaluate(
                    """
                    ({selector, index}) => {
                      const slides = Array.from(document.querySelectorAll(selector));
                      slides.forEach((s, i) => {
                        s.classList.toggle('active', i === index);
                        s.style.display = i === index ? 'flex' : 'none';
                        s.style.opacity = i === index ? '1' : '0';
                        s.style.position = 'relative';
                        s.style.left = '0';
                        s.style.top = '0';
                      });
                    }
                    """,
                    {"selector": selected_candidate, "index": index},
                )
                await page.wait_for_timeout(120)
                png = await page.locator(selected_candidate).nth(index).screenshot(type="png")
                images.append(png)

            await browser.close()
    except Exception as exc:
        issues.append(f"fidelity_render_failed_or_browser_missing:{exc}")

    return images, [_truncate_issue(item) for item in issues], selector_used


async def _render_fidelity_png_units(
    slide_html_units: list[str],
    *,
    shared_styles: str = "",
    width_px: int = 1600,
    height_px: int = 900,
) -> tuple[list[bytes], list[str]]:
    issues: list[str] = []
    images: list[bytes] = []

    try:
        from playwright.async_api import async_playwright
    except Exception as exc:
        return [], [f"fidelity_unavailable_playwright_missing:{exc}"]

    try:
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=True)
            page = await browser.new_page(viewport={"width": width_px, "height": height_px})
            for raw_slide in slide_html_units:
                doc = _build_fidelity_slide_document(raw_slide, shared_styles=shared_styles)
                await page.set_content(doc, wait_until="networkidle", timeout=20000)
                await page.wait_for_timeout(250)
                png = await page.locator("body").screenshot(type="png")
                images.append(png)
            await browser.close()
    except Exception as exc:
        issues.append(f"fidelity_render_failed_or_browser_missing:{exc}")

    return images, [_truncate_issue(item) for item in issues]


async def build_pptx_file_async(request_payload: dict[str, Any], output_path: str) -> dict[str, Any]:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except Exception as exc:
        return build_pptx_error("missing_dependency", f"python-pptx unavailable: {exc}")

    slide_units, slide_units_source, shared_styles = _normalize_slide_html_units(request_payload)
    max_chars_per_slide = int(request_payload.get("paginate", {}).get("max_chars_per_slide") or 900)
    respect_page_break = bool(request_payload.get("paginate", {}).get("respect_page_break", True))
    has_explicit_slides = bool(request_payload.get("slides_html"))

    if has_explicit_slides or len(slide_units) >= 2:
        slides = _slides_from_html_units(
            slide_units,
            max_chars_per_slide=max_chars_per_slide,
            respect_page_break=respect_page_break,
            force_one_slide_per_unit=has_explicit_slides,
        )
    else:
        slides = paginate_html(
            request_payload.get("html") or "",
            mode=str(request_payload.get("paginate", {}).get("mode") or "auto"),
            selector=str(request_payload.get("paginate", {}).get("selector") or ""),
            max_chars_per_slide=max_chars_per_slide,
            respect_page_break=respect_page_break,
        )

    slide_overrides = request_payload.get("paginate", {}).get("slide_overrides")
    if isinstance(slide_overrides, dict):
        slides = apply_slide_overrides(slides, slide_overrides)

    if not slides:
        return build_pptx_error("empty_html", "No usable HTML content after parsing.")

    page = request_payload.get("page", {})
    theme = request_payload.get("theme", {})

    prs = Presentation()
    prs.slide_width = Inches(float(page.get("width_in") or 13.333))
    prs.slide_height = Inches(float(page.get("height_in") or 7.5))

    margin_px = int(page.get("margin_px") or 48)
    margin_in = max(0.1, margin_px / 96.0)

    font_family = str(theme.get("font_family") or "Calibri")
    title_color = _normalize_hex_color(str(theme.get("title_color") or ""), "1F2937")
    text_color = _normalize_hex_color(str(theme.get("text_color") or ""), "111827")

    qa_issues: list[str] = collect_pptx_input_qa_issues(request_payload.get("html") or "")
    render_mode_requested = str(request_payload.get("render_mode") or "auto")
    strict_fidelity = bool(request_payload.get("strict_fidelity", False))
    render_mode_used = "semantic"
    preview_html = _build_html_deck_preview(slide_units, shared_styles=shared_styles)
    preview_height = 560

    selector = str(request_payload.get("paginate", {}).get("selector") or "")
    has_slide_units = bool(slide_units)
    has_multi_slide_preview = has_explicit_slides or len(slide_units) >= 2
    should_try_fidelity = render_mode_requested == "fidelity" or (
        render_mode_requested == "auto" and has_multi_slide_preview
    )
    if should_try_fidelity:
        page_width_in = float(page.get("width_in") or 13.333)
        page_height_in = float(page.get("height_in") or 7.5)
        fidelity_width_px = 1600
        fidelity_height_px = max(720, int(fidelity_width_px * (page_height_in / max(page_width_in, 0.1))))
        png_slides, fidelity_issues = await _render_fidelity_png_units(
            slide_units if has_slide_units else [request_payload.get("html") or ""],
            shared_styles=shared_styles,
            width_px=fidelity_width_px,
            height_px=fidelity_height_px,
        )
        qa_issues.extend(fidelity_issues)
        if png_slides:
            render_mode_used = "fidelity"
            for raw_png in png_slides:
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                slide.shapes.add_picture(
                    io.BytesIO(raw_png),
                    Inches(0),
                    Inches(0),
                    width=prs.slide_width,
                    height=prs.slide_height,
                )
        elif render_mode_requested == "fidelity":
            if strict_fidelity:
                detail = next(
                    (
                        issue.split(":", 1)[1]
                        for issue in fidelity_issues
                        if issue.startswith("fidelity_unavailable_playwright_missing:")
                        or issue.startswith("fidelity_render_failed_or_browser_missing:")
                    ),
                    "",
                )
                message = "High-fidelity export is not available right now."
                if detail:
                    message = f"{message} {detail}"
                return build_pptx_error("fidelity_unavailable", message)
            qa_issues.append("fidelity_requested_but_fallback_to_semantic")
        else:
            qa_issues.append("fidelity_auto_fallback_to_semantic")

    if render_mode_used == "semantic":
        usable_width = float(page.get("width_in") or 13.333) - (margin_in * 2)
        usable_height = float(page.get("height_in") or 7.5) - (margin_in * 2)
        rendered_slides: list[list[dict[str, Any]]] = []

        for source_slide in slides:
            block_queue = [dict(block) for block in source_slide]
            while block_queue:
                slide = prs.slides.add_slide(prs.slide_layouts[6])
                current_slide_blocks: list[dict[str, Any]] = []
                top_in = margin_in
                next_queue: list[dict[str, Any]] = []
                consumed_all = True

                for index, original_block in enumerate(block_queue):
                    block = dict(original_block)
                    force_next_slide_after_render = False
                    block_type = block.get("type")
                    gap_after = 0.08 if block_type in {"heading", "table", "image"} else 0.05
                    required_h = _semantic_block_height(block)
                    available_h = (margin_in + usable_height) - top_in
                    if required_h + gap_after > available_h:
                        head, tail = _split_block_for_available_height(block, available_h - gap_after)
                        if head is None:
                            consumed_all = False
                            next_queue.append(original_block)
                            next_queue.extend(block_queue[index + 1 :])
                            break
                        block = head
                        if tail is not None:
                            next_queue.append(tail)
                        next_queue.extend(block_queue[index + 1 :])
                        consumed_all = False
                        force_next_slide_after_render = True

                    block_type = block.get("type")
                    if block_type == "heading":
                        box_h = 0.8
                        shape = slide.shapes.add_textbox(Inches(margin_in), Inches(top_in), Inches(usable_width), Inches(box_h))
                        tf = shape.text_frame
                        tf.clear()
                        p = tf.paragraphs[0]
                        p.text = str(block.get("text") or "")
                        p.font.bold = True
                        p.font.size = Pt(32 if int(block.get("level") or 2) == 1 else 26)
                        p.font.name = font_family
                        p.font.color.rgb = RGBColor.from_string(title_color)
                        p.alignment = PP_ALIGN.LEFT
                        current_slide_blocks.append(block)
                        top_in += box_h + 0.08
                    elif block_type == "paragraph":
                        text = str(block.get("text") or "")
                        box_h = _estimate_paragraph_height(text)
                        shape = slide.shapes.add_textbox(Inches(margin_in), Inches(top_in), Inches(usable_width), Inches(box_h))
                        tf = shape.text_frame
                        tf.clear()
                        p = tf.paragraphs[0]
                        p.text = text
                        p.font.size = Pt(18)
                        p.font.name = font_family
                        p.font.color.rgb = RGBColor.from_string(text_color)
                        current_slide_blocks.append(block)
                        top_in += box_h + 0.05
                    elif block_type == "bullets":
                        items = [str(i) for i in block.get("items", []) if str(i).strip()]
                        if not items:
                            continue
                        box_h = _estimate_bullets_height(items)
                        shape = slide.shapes.add_textbox(Inches(margin_in), Inches(top_in), Inches(usable_width), Inches(box_h))
                        tf = shape.text_frame
                        tf.clear()
                        ordered = bool(block.get("ordered"))
                        for bullet_index, item in enumerate(items):
                            para = tf.paragraphs[0] if bullet_index == 0 else tf.add_paragraph()
                            para.text = f"{bullet_index + 1}. {item}" if ordered else str(item)
                            para.level = 0
                            para.font.size = Pt(17)
                            para.font.name = font_family
                            para.font.color.rgb = RGBColor.from_string(text_color)
                            para.space_after = Pt(2)
                            if not ordered:
                                para.bullet = True
                        current_slide_blocks.append({"type": "bullets", "ordered": ordered, "items": items})
                        top_in += box_h + 0.05
                    elif block_type == "table":
                        rows = block.get("rows", [])
                        if not rows:
                            continue
                        row_count = len(rows)
                        col_count = max(len(row) for row in rows)
                        box_h = _estimate_table_height(rows)
                        table_shape = slide.shapes.add_table(
                            row_count,
                            col_count,
                            Inches(margin_in),
                            Inches(top_in),
                            Inches(usable_width),
                            Inches(box_h),
                        )
                        table = table_shape.table
                        for r, row in enumerate(rows):
                            for c in range(col_count):
                                table.cell(r, c).text = str(row[c]) if c < len(row) else ""
                        current_slide_blocks.append({"type": "table", "rows": rows})
                        top_in += box_h + 0.08
                    elif block_type == "image":
                        image_data = await _load_image_stream(str(block.get("src") or ""))
                        if image_data is None:
                            continue
                        img_h = min((margin_in + usable_height) - top_in, 2.0)
                        if img_h <= 0.25:
                            continue
                        slide.shapes.add_picture(image_data, Inches(margin_in), Inches(top_in), height=Inches(img_h))
                        current_slide_blocks.append(block)
                        top_in += img_h + 0.08

                    if top_in >= margin_in + usable_height - 0.2:
                        next_queue.extend(block_queue[index + 1 :])
                        consumed_all = False
                        break
                    if force_next_slide_after_render:
                        break

                block_queue = [] if consumed_all else next_queue
                rendered_slides.append(current_slide_blocks or [{"type": "paragraph", "text": "(blank slide)"}])

                # Avoid infinite loops on impossible blocks.
                if not current_slide_blocks and block_queue:
                    qa_issues.append("semantic_overflow_forced_truncate")
                    block_queue = block_queue[1:]

        slides = rendered_slides
        # Keep preview tied to the original HTML/CSS whenever possible.
        # Semantic mode is only the export fallback.
        if not preview_html:
            preview_units = [_build_semantic_slide_doc(blocks, i) for i, blocks in enumerate(slides)]
            preview_html = _build_html_deck_preview(preview_units)

    def _preview_from_slides(slides_data: list[list[dict[str, Any]]]) -> str:
        cards: list[str] = []
        for idx, slide_blocks in enumerate(slides_data, start=1):
            title = f"Slide {idx}"
            summary_parts: list[str] = []
            for block in slide_blocks:
                block_type = block.get("type")
                if block_type == "heading" and title.startswith("Slide"):
                    title = str(block.get("text") or title)
                elif block_type == "paragraph":
                    summary_parts.append(str(block.get("text") or ""))
                elif block_type == "bullets":
                    summary_parts.extend([str(i) for i in block.get("items", [])])
                elif block_type == "table":
                    rows = block.get("rows", [])
                    if rows:
                        summary_parts.append(" / ".join(str(c) for c in rows[0]))
            summary = " ".join(summary_parts).strip()
            if len(summary) > 160:
                summary = summary[:160].rstrip() + "..."
            cards.append(
                "<article style='min-width:220px;max-width:260px;background:#10131a;border:1px solid #2a3140;border-radius:14px;padding:12px;'>"
                f"<div style='font-size:11px;opacity:.7;margin-bottom:6px;'>#{idx}</div>"
                f"<h4 style='margin:0 0 8px;font-size:15px;line-height:1.3;'>{title}</h4>"
                f"<p style='margin:0;font-size:12px;line-height:1.45;opacity:.9;'>{summary}</p>"
                "</article>"
            )
        return (
            "<div style='padding:10px;color:#eaf0ff;font-family:ui-sans-serif,system-ui,sans-serif;'>"
            "<div style='font-size:12px;opacity:.75;margin-bottom:10px;'>Deck Preview</div>"
            "<div style='display:flex;gap:10px;overflow:auto;'>" + "".join(cards) + "</div>"
            "</div>"
        )

    if not preview_html:
        preview_html = _preview_from_slides(slides)

    if render_mode_requested == "fidelity":
        qa_issues.append("fidelity_mode_disabled_use_html_preview")
    if has_explicit_slides:
        qa_issues.append(f"slides_html_mode:{len(slide_units)}")
    elif slide_units_source in {"selector", ".slide", "[data-slide]", "section.slide"} and len(slide_units) >= 2:
        qa_issues.append(f"selector_slides_mode:{len(slide_units)}")

    prs.save(output_path)
    return {
        "type": "pptx_render_result",
        "slide_count": len(prs.slides),
        "preview_html": preview_html,
        "preview_height": preview_height,
        "qa_issues": [_truncate_issue(item) for item in qa_issues],
        "render_mode_used": render_mode_used,
    }


def build_pptx_file(request_payload: dict[str, Any], output_path: str) -> dict[str, Any]:
    return asyncio.run(build_pptx_file_async(request_payload, output_path))
